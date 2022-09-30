import string
from azure.cognitiveservices.vision.computervision import ComputerVisionClient
from msrest.authentication import CognitiveServicesCredentials
from azure.cognitiveservices.vision.computervision.models import OperationStatusCodes
import time
from dotenv import load_dotenv
import os
import docx
from fpdf import FPDF

def main():
    try:
        # Get Configuration Settings
        load_dotenv()
        endpoint = os.getenv("COG_SERVICE_ENDPOINT")
        key = os.getenv("COG_SERVICE_KEY")
        # Authenticate Computer Vision client
        computervision_client = ComputerVisionClient(endpoint, CognitiveServicesCredentials(key))
        # Extract test
        images_folder = os.path.join (os.path.dirname(os.path.abspath(__file__)), "images")
        read_image_path = os.path.join (images_folder, "test4.png")
        get_doc(read_image_path,computervision_client)
        get_pdf(read_image_path,computervision_client)
        
    
    except Exception as ex:
        print(ex)


def get_text(image_file,computervision_client):
    
    with open(image_file, "rb") as image:
        read_response = computervision_client.read_in_stream(image, raw=True)

    read_operation_location = read_response.headers["Operation-Location"]
    operation_id = read_operation_location.split("/")[-1]
    filename = os.path.splitext(image_file)[0]

    while True:
        read_result = computervision_client.get_read_result(operation_id)
        if read_result.status.lower() not in ['notstarted', 'running']:
            break
        time.sleep(1)
    
    with open(filename + '.txt', 'w') as f:
        if read_result.status == OperationStatusCodes.succeeded:
            for page in read_result.analyze_result.read_results:
                for line in page.lines:
                    f.write(line.text + "\n")

def get_doc(image_file, computervision_client):
    get_text(image_file, computervision_client)
    filename = os.path.splitext(image_file)[0]
    doc = docx.Document()
    with open(filename + ".txt", 'r', encoding='utf-8') as openfile:
        line = openfile.read()
        doc.add_paragraph(line)
        doc.save(filename + ".docx")

def get_pdf(image_file, computervision_client):
    get_text(image_file, computervision_client)
    filename = os.path.splitext(image_file)[0]
    pdf=FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size = 15)
    with open(filename + ".txt", 'r', encoding='utf-8') as openfile:
        line = openfile.read()
        pdf.cell(200, 10, txt = line,
            ln = 1, align = 'C')

    pdf.output(filename+".pdf")              


if __name__ == "__main__":
    main()